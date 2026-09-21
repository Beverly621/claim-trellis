from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from claim_trellis.citations import extract_citation_sentences
from claim_trellis.models import ParsedDocument


class IngestionError(ValueError):
    pass


SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> tuple[str, list[str]]:
    if not data.startswith(b"%PDF"):
        raise IngestionError("The file extension is PDF but the signature is not a PDF.")
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # pypdf raises several parser-specific exception classes
        raise IngestionError(f"PDF parsing failed: {exc}") from exc
    text = "\n\n".join(pages)
    warnings: list[str] = []
    if not text.strip():
        warnings.append("No selectable text was found; scanned PDFs require OCR before upload.")
    return text, warnings


def _read_docx(data: bytes) -> tuple[str, list[str]]:
    if not data.startswith(b"PK"):
        raise IngestionError(
            "The file extension is DOCX but the signature is not an Office archive."
        )
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"DOCX parsing failed: {exc}") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n\n".join(blocks), []


def parse_document_bytes(
    filename: str,
    data: bytes,
    media_type: str = "application/octet-stream",
    *,
    max_bytes: int = 25 * 1024 * 1024,
    max_chars: int = 2_000_000,
) -> ParsedDocument:
    if len(data) > max_bytes:
        raise IngestionError(f"Upload exceeds the {max_bytes:,}-byte limit.")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise IngestionError(
            f"Unsupported file type {suffix or '(none)'}. Use TXT, MD, PDF, or DOCX."
        )

    warnings: list[str] = []
    if suffix == ".pdf":
        raw_text, warnings = _read_pdf(data)
    elif suffix == ".docx":
        raw_text, warnings = _read_docx(data)
    else:
        try:
            raw_text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestionError("Text files must use UTF-8 encoding.") from exc

    text = _clean_text(raw_text)
    if len(text) > max_chars:
        raise IngestionError(f"Parsed text exceeds the {max_chars:,}-character limit.")
    if not text:
        raise IngestionError("The document contains no readable text.")

    return ParsedDocument(
        filename=Path(filename).name,
        media_type=media_type,
        text=text,
        content_sha256=hashlib.sha256(text.encode()).hexdigest(),
        character_count=len(text),
        citation_sentences=extract_citation_sentences(text),
        warnings=warnings,
    )
